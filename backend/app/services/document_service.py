"""
Document Service

문서 업로드, 저장, OCR 처리 비즈니스 로직
"""

import io
import logging
import mimetypes
from datetime import datetime, timezone
from pathlib import Path
from typing import BinaryIO
from uuid import UUID, uuid4

from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.core.exceptions import (
    DocumentNotFoundError,
    DocumentProcessingError,
    FileSizeExceededError,
    InvalidFileTypeError,
    StorageQuotaExceededError,
    UnauthorizedError,
)
from app.models.document import Document
from app.models.user import User
from app.repositories.document import DocumentRepository

logger = logging.getLogger(__name__)

# File type configurations
ALLOWED_EXTENSIONS = {
    "pdf": ["application/pdf"],
    "image": ["image/jpeg", "image/jpg", "image/png", "image/gif", "image/webp"],
    "text": ["text/plain"],
    "word": [
        "application/msword",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ],
    "excel": [
        "application/vnd.ms-excel",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ],
}

# Flatten for quick lookup
ALLOWED_MIME_TYPES = {
    mime: file_type
    for file_type, mimes in ALLOWED_EXTENSIONS.items()
    for mime in mimes
}

# Size limits by tier (in bytes)
FILE_SIZE_LIMITS = {
    "basic": 5 * 1024 * 1024,  # 5MB
    "pro": 25 * 1024 * 1024,  # 25MB
    "enterprise": 100 * 1024 * 1024,  # 100MB
}

# Storage quota by tier (in bytes)
STORAGE_QUOTAS = {
    "basic": 100 * 1024 * 1024,  # 100MB
    "pro": 1 * 1024 * 1024 * 1024,  # 1GB
    "enterprise": 10 * 1024 * 1024 * 1024,  # 10GB
}


class DocumentService:
    """Service for managing user documents."""

    def __init__(self, db: AsyncSession):
        self.db = db
        self.document_repo = DocumentRepository(db)
        self._s3_client = None

    # =========================================================================
    # S3 Client Management
    # =========================================================================

    async def _get_s3_client(self):
        """Get or create S3 client."""
        if self._s3_client is None:
            try:
                import aioboto3

                session = aioboto3.Session()
                self._s3_client = await session.client(
                    "s3",
                    region_name=settings.AWS_REGION,
                    aws_access_key_id=settings.AWS_ACCESS_KEY_ID,
                    aws_secret_access_key=settings.AWS_SECRET_ACCESS_KEY,
                ).__aenter__()
            except ImportError:
                logger.warning("aioboto3 not installed, S3 operations will fail")
                raise DocumentProcessingError("Storage service not available")
            except Exception as e:
                logger.error(f"Failed to initialize S3 client: {e}")
                raise DocumentProcessingError("Storage service initialization failed")

        return self._s3_client

    async def close(self):
        """Close S3 client connection."""
        if self._s3_client:
            await self._s3_client.__aexit__(None, None, None)
            self._s3_client = None

    # =========================================================================
    # File Validation
    # =========================================================================

    def _validate_file_type(
        self,
        filename: str,
        content_type: str | None,
    ) -> tuple[str, str]:
        """
        Validate file type and return (file_type, mime_type).

        Raises:
            InvalidFileTypeError: If file type is not allowed
        """
        # Get mime type from content_type or guess from filename
        mime_type = content_type
        if not mime_type or mime_type == "application/octet-stream":
            mime_type, _ = mimetypes.guess_type(filename)

        if not mime_type:
            raise InvalidFileTypeError(
                f"파일 타입을 확인할 수 없습니다: {filename}"
            )

        file_type = ALLOWED_MIME_TYPES.get(mime_type)
        if not file_type:
            raise InvalidFileTypeError(
                f"지원하지 않는 파일 형식입니다: {mime_type}. "
                f"지원 형식: PDF, 이미지(JPG, PNG), Word, Excel, 텍스트"
            )

        return file_type, mime_type

    async def _validate_file_size(
        self,
        file_size: int,
        user: User,
    ) -> None:
        """
        Validate file size against user tier limits.

        Raises:
            FileSizeExceededError: If file exceeds size limit
        """
        tier = user.tier.value if user.tier else "basic"
        limit = FILE_SIZE_LIMITS.get(tier, FILE_SIZE_LIMITS["basic"])

        if file_size > limit:
            limit_mb = limit / (1024 * 1024)
            raise FileSizeExceededError(
                f"파일 크기가 제한을 초과했습니다. "
                f"최대 {limit_mb:.0f}MB까지 업로드 가능합니다."
            )

    async def _validate_storage_quota(
        self,
        additional_size: int,
        user: User,
    ) -> None:
        """
        Validate user's storage quota.

        Raises:
            StorageQuotaExceededError: If quota would be exceeded
        """
        tier = user.tier.value if user.tier else "basic"
        quota = STORAGE_QUOTAS.get(tier, STORAGE_QUOTAS["basic"])

        current_usage = await self.document_repo.get_storage_usage(user.id)
        if current_usage + additional_size > quota:
            quota_mb = quota / (1024 * 1024)
            usage_mb = current_usage / (1024 * 1024)
            raise StorageQuotaExceededError(
                f"저장 공간이 부족합니다. "
                f"현재 사용량: {usage_mb:.1f}MB / {quota_mb:.0f}MB"
            )

    # =========================================================================
    # Document Operations
    # =========================================================================

    async def upload_document(
        self,
        user: User,
        filename: str,
        content: bytes,
        content_type: str | None = None,
        description: str | None = None,
    ) -> Document:
        """
        Upload and store a document.

        Args:
            user: User uploading the document
            filename: Original filename
            content: File content as bytes
            content_type: MIME type of the file
            description: Optional description

        Returns:
            Created Document model

        Raises:
            InvalidFileTypeError: If file type not allowed
            FileSizeExceededError: If file too large
            StorageQuotaExceededError: If storage quota exceeded
            DocumentProcessingError: If upload fails
        """
        # Validate file
        file_type, mime_type = self._validate_file_type(filename, content_type)
        file_size = len(content)
        await self._validate_file_size(file_size, user)
        await self._validate_storage_quota(file_size, user)

        # Generate storage path
        file_ext = Path(filename).suffix.lower()
        unique_filename = f"{uuid4()}{file_ext}"
        storage_path = f"documents/{user.id}/{unique_filename}"

        try:
            # Upload to S3
            s3 = await self._get_s3_client()
            await s3.put_object(
                Bucket=settings.AWS_S3_BUCKET,
                Key=storage_path,
                Body=content,
                ContentType=mime_type,
                Metadata={
                    "original_name": filename,
                    "user_id": str(user.id),
                },
            )

            logger.info(f"Document uploaded to S3: {storage_path}")

        except DocumentProcessingError:
            raise
        except Exception as e:
            logger.error(f"Failed to upload document to S3: {e}")
            raise DocumentProcessingError("문서 업로드에 실패했습니다.")

        # Create database record
        document = await self.document_repo.create(
            user_id=user.id,
            file_name=unique_filename,
            original_name=filename,
            file_type=file_type,
            mime_type=mime_type,
            file_size=file_size,
            storage_path=storage_path,
            storage_bucket=settings.AWS_S3_BUCKET,
            description=description,
            is_processed=False,
            ocr_completed=False,
        )

        logger.info(f"Document record created: {document.id}")
        return document

    async def get_document(
        self,
        document_id: UUID,
        user_id: UUID,
    ) -> Document | None:
        """Get document with ownership check."""
        document = await self.document_repo.get(document_id)

        if not document:
            return None

        if document.is_deleted:
            return None

        if document.user_id != user_id:
            raise UnauthorizedError("이 문서에 접근할 권한이 없습니다.")

        return document

    async def list_documents(
        self,
        user_id: UUID,
        page: int = 1,
        page_size: int = 20,
        file_type: str | None = None,
    ) -> dict:
        """List user's documents with pagination."""
        skip = (page - 1) * page_size

        documents = await self.document_repo.get_user_documents(
            user_id=user_id,
            skip=skip,
            limit=page_size,
            file_type=file_type,
        )

        total = await self.document_repo.count_user_documents(user_id)
        total_pages = (total + page_size - 1) // page_size

        return {
            "items": documents,
            "total": total,
            "page": page,
            "page_size": page_size,
            "total_pages": total_pages,
        }

    async def delete_document(
        self,
        document_id: UUID,
        user_id: UUID,
    ) -> None:
        """Soft delete a document."""
        document = await self.get_document(document_id, user_id)

        if not document:
            raise DocumentNotFoundError("문서를 찾을 수 없습니다.")

        # Soft delete (keeps S3 file for potential recovery)
        await self.document_repo.soft_delete(document_id)

        logger.info(f"Document soft deleted: {document_id}")

    async def permanently_delete_document(
        self,
        document_id: UUID,
        user_id: UUID,
    ) -> None:
        """Permanently delete document from S3 and database."""
        document = await self.document_repo.get(document_id)

        if not document:
            raise DocumentNotFoundError("문서를 찾을 수 없습니다.")

        if document.user_id != user_id:
            raise UnauthorizedError("이 문서에 접근할 권한이 없습니다.")

        # Delete from S3
        try:
            s3 = await self._get_s3_client()
            await s3.delete_object(
                Bucket=document.storage_bucket,
                Key=document.storage_path,
            )
            logger.info(f"Document deleted from S3: {document.storage_path}")
        except Exception as e:
            logger.error(f"Failed to delete from S3: {e}")
            # Continue with database deletion even if S3 fails

        # Delete from database
        await self.document_repo.delete(document_id)
        logger.info(f"Document permanently deleted: {document_id}")

    async def download_document(
        self,
        document_id: UUID,
        user_id: UUID,
    ) -> tuple[bytes, str, str]:
        """
        Download document content.

        Returns:
            Tuple of (content_bytes, filename, content_type)
        """
        document = await self.get_document(document_id, user_id)

        if not document:
            raise DocumentNotFoundError("문서를 찾을 수 없습니다.")

        try:
            s3 = await self._get_s3_client()
            response = await s3.get_object(
                Bucket=document.storage_bucket,
                Key=document.storage_path,
            )
            content = await response["Body"].read()

            return content, document.original_name, document.mime_type

        except Exception as e:
            logger.error(f"Failed to download document: {e}")
            raise DocumentProcessingError("문서 다운로드에 실패했습니다.")

    async def get_presigned_url(
        self,
        document_id: UUID,
        user_id: UUID,
        expires_in: int = 3600,
    ) -> str:
        """
        Generate presigned URL for document download.

        Args:
            document_id: Document ID
            user_id: User ID for ownership check
            expires_in: URL expiration in seconds (default 1 hour)

        Returns:
            Presigned URL string
        """
        document = await self.get_document(document_id, user_id)

        if not document:
            raise DocumentNotFoundError("문서를 찾을 수 없습니다.")

        try:
            s3 = await self._get_s3_client()
            url = await s3.generate_presigned_url(
                "get_object",
                Params={
                    "Bucket": document.storage_bucket,
                    "Key": document.storage_path,
                    "ResponseContentDisposition": f'attachment; filename="{document.original_name}"',
                },
                ExpiresIn=expires_in,
            )
            return url

        except Exception as e:
            logger.error(f"Failed to generate presigned URL: {e}")
            raise DocumentProcessingError("다운로드 링크 생성에 실패했습니다.")

    # =========================================================================
    # OCR Processing
    # =========================================================================

    async def process_document_ocr(
        self,
        document_id: UUID,
    ) -> Document:
        """
        Process document with OCR to extract text.

        This is typically called as a background task after upload.
        """
        document = await self.document_repo.get(document_id)

        if not document:
            raise DocumentNotFoundError("문서를 찾을 수 없습니다.")

        if document.ocr_completed:
            return document

        try:
            # Download document content
            s3 = await self._get_s3_client()
            response = await s3.get_object(
                Bucket=document.storage_bucket,
                Key=document.storage_path,
            )
            content = await response["Body"].read()

            # Extract text based on file type
            extracted_text = ""
            page_count = None
            confidence = None

            if document.file_type == "pdf":
                extracted_text, page_count, confidence = await self._extract_pdf_text(
                    content
                )
            elif document.file_type == "image":
                extracted_text, confidence = await self._extract_image_text(content)
            elif document.file_type == "text":
                extracted_text = content.decode("utf-8", errors="ignore")
            elif document.file_type in ("word", "excel"):
                extracted_text = await self._extract_office_text(
                    content, document.mime_type
                )

            # Update document with OCR results
            document = await self.document_repo.mark_as_processed(
                document_id,
                extracted_text=extracted_text,
                ocr_confidence=confidence,
                page_count=page_count,
            )

            logger.info(f"OCR completed for document: {document_id}")
            return document

        except Exception as e:
            logger.error(f"OCR processing failed for document {document_id}: {e}")
            # Mark as processed with error
            await self.document_repo.update(
                document_id,
                is_processed=True,
                ocr_completed=False,
            )
            raise DocumentProcessingError(f"OCR 처리에 실패했습니다: {str(e)}")

    async def _extract_pdf_text(
        self,
        content: bytes,
    ) -> tuple[str, int, float | None]:
        """Extract text from PDF using PyMuPDF."""
        try:
            import fitz  # PyMuPDF

            pdf = fitz.open(stream=content, filetype="pdf")
            texts = []
            page_count = len(pdf)

            for page in pdf:
                text = page.get_text()
                texts.append(text)

            pdf.close()

            # If no text found, try OCR
            full_text = "\n\n".join(texts)
            if not full_text.strip():
                # Use Tesseract for image-based PDFs
                full_text, confidence = await self._ocr_pdf_pages(content)
                return full_text, page_count, confidence

            return full_text, page_count, None

        except ImportError:
            logger.warning("PyMuPDF not installed, PDF extraction unavailable")
            return "", 0, None
        except Exception as e:
            logger.error(f"PDF extraction failed: {e}")
            raise

    async def _ocr_pdf_pages(
        self,
        content: bytes,
    ) -> tuple[str, float]:
        """OCR PDF pages using Tesseract."""
        try:
            import fitz
            from PIL import Image
            import pytesseract

            pdf = fitz.open(stream=content, filetype="pdf")
            texts = []
            confidences = []

            for page in pdf:
                # Convert page to image
                pix = page.get_pixmap(dpi=300)
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)

                # OCR the image
                data = pytesseract.image_to_data(
                    img, lang="kor+eng", output_type=pytesseract.Output.DICT
                )
                text = " ".join(data["text"])
                texts.append(text)

                # Calculate confidence
                conf_values = [c for c in data["conf"] if c > 0]
                if conf_values:
                    confidences.append(sum(conf_values) / len(conf_values))

            pdf.close()

            full_text = "\n\n".join(texts)
            avg_confidence = sum(confidences) / len(confidences) if confidences else 0

            return full_text, avg_confidence / 100  # Convert to 0-1 scale

        except ImportError:
            logger.warning("Tesseract/PIL not installed, PDF OCR unavailable")
            return "", 0.0
        except Exception as e:
            logger.error(f"PDF OCR failed: {e}")
            raise

    async def _extract_image_text(
        self,
        content: bytes,
    ) -> tuple[str, float]:
        """Extract text from image using Tesseract."""
        try:
            from PIL import Image
            import pytesseract

            img = Image.open(io.BytesIO(content))

            # OCR the image
            data = pytesseract.image_to_data(
                img, lang="kor+eng", output_type=pytesseract.Output.DICT
            )
            text = " ".join(data["text"])

            # Calculate confidence
            conf_values = [c for c in data["conf"] if c > 0]
            avg_confidence = sum(conf_values) / len(conf_values) if conf_values else 0

            return text, avg_confidence / 100

        except ImportError:
            logger.warning("Tesseract/PIL not installed, image OCR unavailable")
            return "", 0.0
        except Exception as e:
            logger.error(f"Image OCR failed: {e}")
            raise

    async def _extract_office_text(
        self,
        content: bytes,
        mime_type: str,
    ) -> str:
        """Extract text from Word/Excel documents."""
        try:
            if "word" in mime_type or mime_type.endswith("document"):
                from docx import Document as DocxDocument

                doc = DocxDocument(io.BytesIO(content))
                texts = [para.text for para in doc.paragraphs]
                return "\n".join(texts)

            elif "excel" in mime_type or mime_type.endswith("sheet"):
                import openpyxl

                wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
                texts = []

                for sheet in wb.worksheets:
                    for row in sheet.iter_rows():
                        row_texts = [str(cell.value) for cell in row if cell.value]
                        if row_texts:
                            texts.append(" | ".join(row_texts))

                wb.close()
                return "\n".join(texts)

            return ""

        except ImportError as e:
            logger.warning(f"Office document library not installed: {e}")
            return ""
        except Exception as e:
            logger.error(f"Office document extraction failed: {e}")
            raise

    # =========================================================================
    # OCR Result Access
    # =========================================================================

    async def get_ocr_result(
        self,
        document_id: UUID,
        user_id: UUID,
    ) -> dict:
        """Get OCR extraction result for a document."""
        document = await self.get_document(document_id, user_id)

        if not document:
            raise DocumentNotFoundError("문서를 찾을 수 없습니다.")

        if not document.is_processed:
            return {
                "document_id": str(document_id),
                "status": "processing",
                "extracted_text": None,
                "confidence": None,
                "page_count": None,
            }

        if not document.ocr_completed:
            return {
                "document_id": str(document_id),
                "status": "failed",
                "extracted_text": None,
                "confidence": None,
                "page_count": None,
            }

        return {
            "document_id": str(document_id),
            "status": "completed",
            "extracted_text": document.extracted_text,
            "confidence": document.ocr_confidence,
            "page_count": document.page_count,
        }

    # =========================================================================
    # Storage Statistics
    # =========================================================================

    async def get_storage_stats(self, user_id: UUID, user_tier: str) -> dict:
        """Get storage usage statistics for a user."""
        current_usage = await self.document_repo.get_storage_usage(user_id)
        quota = STORAGE_QUOTAS.get(user_tier, STORAGE_QUOTAS["basic"])
        document_count = await self.document_repo.count_user_documents(user_id)

        return {
            "used_bytes": current_usage,
            "quota_bytes": quota,
            "used_percentage": (current_usage / quota * 100) if quota > 0 else 0,
            "document_count": document_count,
            "file_size_limit": FILE_SIZE_LIMITS.get(user_tier, FILE_SIZE_LIMITS["basic"]),
        }
